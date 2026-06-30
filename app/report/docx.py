"""DOCX 리포트 (python-docx) — 웹 대시보드 레이아웃을 따른 편집 가능한 Word 문서.

구성(대시보드와 동일 흐름):
    헤더 → 종합 평가(레이더+카테고리 점수) → 분석 요약 → 카테고리별 상세(항목표: 코멘트+근거)
    → 최종 피드백(강점/개선) → (다강의 시) 점수 추이

LLM 미사용 — 이미 분석된 스코어카드(항목별 reason/근거/강점/개선)를 렌더링만 한다.
"""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.analysis.rubrics import load_item_rubrics
from app.analysis.schemas import InstructorScorecard
from app.report import charts, html

# ── 대시보드와 동일한 카테고리 한글 라벨/순서 ──────────────────────────────
CATEGORY_LABEL = {
    "structure": "강의 도입·구조",
    "concept": "개념 명확성",
    "practice": "예시·실습 연계",
    "language": "언어 표현 품질",
    "interaction": "수강생 상호작용",
}
CATEGORY_ORDER = ["structure", "concept", "practice", "language", "interaction"]

# 색상 (대시보드 톤)
_HEAD_FILL = "EEF2FF"  # 연한 인디고 (표 헤더)
_ZEBRA = "F8FAFC"      # 행 줄무늬
_GREEN = RGBColor(0x05, 0x96, 0x69)
_AMBER = RGBColor(0xD9, 0x77, 0x06)
_RED = RGBColor(0xDC, 0x26, 0x26)
_SUBTLE = RGBColor(0x6B, 0x72, 0x80)
_INK = RGBColor(0x37, 0x41, 0x51)


def _score_color(score: float) -> RGBColor:
    if score >= 4.0:
        return _GREEN
    if score >= 3.0:
        return _AMBER
    return _RED


def _band(score: float) -> str:
    if score >= 4.5:
        return "우수"
    if score >= 3.5:
        return "양호"
    if score >= 2.5:
        return "보통"
    if score >= 1.5:
        return "미흡"
    return "부족"


def _shade(cell, hex_fill: str) -> None:
    """셀 배경색."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_widths(table, widths_in: list[float]) -> None:
    """열 너비(inch) 고정."""
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def _header_row(table, labels: list[str]) -> None:
    for cell, label in zip(table.rows[0].cells, labels):
        _shade(cell, _HEAD_FILL)
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = _INK


def _set_base_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "맑은 고딕")


def render(
    scorecard: InstructorScorecard,
    output_path: str,
    radar_png: bytes | None = None,
    trend_png: bytes | None = None,
) -> str:
    """Scorecard → DOCX 파일 경로."""
    radar_png = radar_png or charts.radar_chart(scorecard)
    trend_png = trend_png or charts.trend_chart(scorecard)
    rubrics = load_item_rubrics()

    doc = Document()
    _set_base_font(doc)

    # ── 헤더 ─────────────────────────────────────────────────
    title = doc.add_heading("강의 분석 리포트", level=0)
    title.runs[0].font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    meta = doc.add_paragraph()
    meta.add_run(f"강사 {scorecard.instructor_id}").bold = True
    meta.add_run(f"   ·   강의일 {scorecard.lecture_date}").font.color.rgb = _SUBTLE
    meta.add_run(
        f"   ·   생성 {scorecard.analyzed_at.strftime('%Y-%m-%d %H:%M')}"
    ).font.color.rgb = _SUBTLE

    # ── 종합 평가 ────────────────────────────────────────────
    doc.add_heading("종합 평가", level=1)
    score_p = doc.add_paragraph()
    big = score_p.add_run(f"{scorecard.overall_score:.1f}")
    big.font.size = Pt(30)
    big.bold = True
    big.font.color.rgb = _score_color(scorecard.overall_score)
    score_p.add_run(" / 5.0   ").font.size = Pt(14)
    badge = score_p.add_run(_band(scorecard.overall_score))
    badge.bold = True
    badge.font.color.rgb = _score_color(scorecard.overall_score)
    if scorecard.trend_label:
        tp = doc.add_paragraph()
        tp.add_run(
            f"추이: {scorecard.trend_label} (기울기 {scorecard.trend_slope or 0:.2f})"
        ).font.color.rgb = _SUBTLE

    doc.add_picture(io.BytesIO(radar_png), width=Inches(4.0))

    # 카테고리 점수 표 (점수 막대 포함)
    cat_tbl = doc.add_table(rows=1, cols=4)
    cat_tbl.style = "Table Grid"
    cat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    _header_row(cat_tbl, ["카테고리", "점수", "막대", "가중치"])
    by_cat = {c.category: c for c in scorecard.category_scores}
    for i, cat in enumerate(CATEGORY_ORDER):
        c = by_cat.get(cat)
        if not c:
            continue
        cells = cat_tbl.add_row().cells
        if i % 2 == 1:
            for cell in cells:
                _shade(cell, _ZEBRA)
        cells[0].paragraphs[0].add_run(CATEGORY_LABEL.get(cat, cat))
        r = cells[1].paragraphs[0].add_run(f"{c.score:.1f}")
        r.bold = True
        r.font.color.rgb = _score_color(c.score)
        n = max(0, min(5, round(c.score)))
        br = cells[2].paragraphs[0].add_run("█" * n + "░" * (5 - n))
        br.font.color.rgb = _score_color(c.score)
        cells[3].paragraphs[0].add_run(f"{c.weight * 100:.0f}%").font.color.rgb = _SUBTLE
    _set_widths(cat_tbl, [2.0, 0.8, 1.6, 1.0])

    # ── 분석 요약 (점수 기반 템플릿, LLM 미사용) ──────────────
    doc.add_heading("분석 요약", level=1)
    cats_sorted = sorted(scorecard.category_scores, key=lambda c: c.score, reverse=True)
    if cats_sorted:
        best, worst = cats_sorted[0], cats_sorted[-1]
        doc.add_paragraph(
            f"종합 {scorecard.overall_score:.1f} / 5.0 — {_band(scorecard.overall_score)} 수준입니다. "
            f"가장 강한 영역은 ‘{CATEGORY_LABEL.get(best.category, best.category)}’"
            f"({best.score:.1f}점), 가장 약한 영역은 ‘{CATEGORY_LABEL.get(worst.category, worst.category)}’"
            f"({worst.score:.1f}점)입니다.",
            style="List Bullet",
        )
    lows = sorted(scorecard.item_scores, key=lambda it: it.final_score)[:3]
    if lows:
        doc.add_paragraph(
            "우선 점검이 필요한 항목: "
            + ", ".join(f"{it.name}({it.final_score:.1f})" for it in lows)
            + ".",
            style="List Bullet",
        )

    # ── 카테고리별 상세 (항목표: 코멘트 + 근거) ───────────────
    doc.add_heading("상세 설명", level=1)
    for cat in CATEGORY_ORDER:
        items = sorted(
            [it for it in scorecard.item_scores if it.category == cat],
            key=lambda it: it.item_id,
        )
        if not items:
            continue
        h = doc.add_heading(CATEGORY_LABEL.get(cat, cat), level=2)
        h.runs[0].font.color.rgb = _INK

        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        _header_row(tbl, ["항목", "항목 설명", "점수", "해설"])
        for i, it in enumerate(items):
            cells = tbl.add_row().cells
            if i % 2 == 1:
                for cell in cells:
                    _shade(cell, _ZEBRA)
            # 항목 (id + 이름)
            nm = cells[0].paragraphs[0]
            nm.add_run(f"{it.item_id} ").font.color.rgb = _SUBTLE
            nm.add_run(it.name).bold = True
            if it.needs_human_review:
                rv = cells[0].add_paragraph().add_run("검토 권장")
                rv.font.size = Pt(8)
                rv.font.color.rgb = _AMBER
            # 항목 설명 (루브릭 세부기준)
            crit = (rubrics.get(it.item_id) or {}).get("criterion", "")
            cd = cells[1].paragraphs[0].add_run(crit)
            cd.font.size = Pt(9)
            cd.font.color.rgb = _SUBTLE
            # 점수
            sr = cells[2].paragraphs[0].add_run(f"{it.final_score:.1f}")
            sr.bold = True
            sr.font.color.rgb = _score_color(it.final_score)
            # 해설 (코멘트 + 근거)
            comment = it.reason or it.improvements or ""
            cells[3].paragraphs[0].add_run(comment).font.size = Pt(9)
            if it.evidence:
                gp = cells[3].add_paragraph()
                lbl = gp.add_run("근거 ")
                lbl.bold = True
                lbl.font.size = Pt(8)
                lbl.font.color.rgb = _SUBTLE
                gv = gp.add_run(it.evidence)
                gv.font.size = Pt(8)
                gv.font.color.rgb = _SUBTLE
        _set_widths(tbl, [1.4, 2.2, 0.6, 3.0])

    # ── 최종 피드백 (강점 / 개선) ─────────────────────────────
    doc.add_heading("최종 피드백", level=1)
    sh = doc.add_paragraph().add_run("잘하고 있는 점")
    sh.bold = True
    sh.font.color.rgb = _GREEN
    for it in html.top_items(scorecard, best=True):
        doc.add_paragraph(
            f"{it.name} ({it.final_score:.1f}) — {it.strengths or it.reason or ''}",
            style="List Bullet",
        )
    wh = doc.add_paragraph().add_run("우선 개선할 점")
    wh.bold = True
    wh.font.color.rgb = _RED
    for it in html.top_items(scorecard, best=False):
        doc.add_paragraph(
            f"{it.name} ({it.final_score:.1f}) — {it.improvements or it.reason or ''}",
            style="List Bullet",
        )

    # ── 점수 추이 (다강의 시) ─────────────────────────────────
    if scorecard.trend_points and len(scorecard.trend_points) >= 2:
        doc.add_heading("점수 추이", level=1)
        doc.add_picture(io.BytesIO(trend_png), width=Inches(5.5))

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return str(out)
